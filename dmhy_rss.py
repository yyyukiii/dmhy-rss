# -*- coding: utf-8 -*-
"""
DMHY (动漫花园) RSS 生成器
功能：抓取动漫花园资源列表全部页 + 详情页正文 -> 增量去重 -> 生成 RSS 2.0 + xlsx
用法：
  python dmhy_rss.py                 # 增量抓取（新条目抓详情正文）
  python dmhy_rss.py --force         # 忽略去重状态，重新抓取并生成全部
依赖：requests beautifulsoup4 feedgen openpyxl
"""
import argparse
import json
import os
import re
import sys
import time
from datetime import datetime, timedelta, timezone

import requests
from bs4 import BeautifulSoup
from feedgen.feed import FeedGenerator

# ---------------------------------------------------------------------------
# 配置（可按需修改）
# ---------------------------------------------------------------------------
BASE_URL = "https://share.dmhy.org"
LIST_BASE = "https://share.dmhy.org/topics/list"
QUERY = "keyword=&sort_id=0&team_id=581&order=date-desc"
MAX_PAGES = 500          # 列表翻页安全上限（实际自动停在最后一页）
PAGE_DELAY = 0.3         # 列表页之间间隔
DETAIL_DELAY = 0.2       # 详情页之间间隔
DETAIL_BATCH = 400       # 每次运行最多补爬多少条无正文的历史详情
FETCH_DETAIL = True      # 是否抓取详情页正文
FEED_TITLE = "VCB-Studio - 动漫花园资源网"
FEED_DESC = "VCB-Studio 字幕组在动漫花园的发布（含简介正文，本地自建 RSS）"
FEED_LINK = f"{LIST_BASE}?{QUERY}"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
}

DEFAULT_OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dmhy_feed.xml")
XLSX_OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dmhy_titles.xlsx")
STATE_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dmhy_state.json")

TZ_CST = timezone(timedelta(hours=8), "CST")
DATE_RE = re.compile(r"(\d{4})/(\d{2})/(\d{2}) (\d{2}):(\d{2})")
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def safe_xml(s) -> str:
    """去掉 XML 1.0 不允许的控制字符（保留 \\t \\n \\r）。"""
    if not s:
        return ""
    return _CONTROL_RE.sub("", str(s))


def page_url(page: int) -> str:
    if page <= 1:
        return f"{LIST_BASE}?{QUERY}"
    return f"{LIST_BASE}/page/{page}?{QUERY}"


def fetch_html(url: str) -> str:
    resp = requests.get(url, headers=HEADERS, timeout=30)
    resp.raise_for_status()
    resp.encoding = "utf-8"
    return resp.text


def has_next_page(soup: BeautifulSoup) -> bool:
    for a in soup.find_all("a", href=re.compile(r"/topics/list/page/")):
        if "下一頁" in a.get_text():
            return True
    return False


def parse_list_page(html: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find("table", id="topic_list")
    if table is None:
        return []
    topics = []
    for tr in table.select("tbody tr"):
        tds = tr.find_all("td")
        if len(tds) < 6:
            continue
        pub_dt = None
        hidden = tds[0].find("span", style=lambda v: v and "display" in v)
        if hidden:
            m = DATE_RE.search(hidden.get_text())
            if m:
                y, mo, d, h, mi = map(int, m.groups())
                pub_dt = datetime(y, mo, d, h, mi, tzinfo=TZ_CST)
        title_a = tds[2].find("a", href=re.compile(r"/topics/view/"))
        if title_a is None:
            continue
        title = title_a.get_text(strip=True)
        href = title_a.get("href", "")
        topic_url = href if href.startswith("http") else BASE_URL + href
        m_id = re.search(r"/topics/view/(\d+)", href)
        topic_id = m_id.group(1) if m_id else topic_url
        cat = tds[1].get_text(strip=True) if len(tds) > 1 else ""
        magnet = ""
        m_a = tr.find("a", class_="arrow-magnet")
        if m_a is not None:
            magnet = m_a.get("href", "")
        size = tds[4].get_text(strip=True) if len(tds) > 4 else ""
        publisher = tds[-1].get_text(strip=True) if len(tds) > 6 else ""
        topics.append({
            "id": topic_id, "title": title, "link": topic_url,
            "pub_date": pub_dt.isoformat() if pub_dt else None,
            "category": cat, "magnet": magnet, "size": size,
            "publisher": publisher, "content": "",
        })
    return topics


def fetch_detail_content(url: str) -> str:
    """抓取详情页简介正文（保留 HTML 以保留图片和链接）。"""
    try:
        html = fetch_html(url)
    except Exception as e:
        print(f"      !! 详情页抓取失败 {url}: {e}")
        return ""
    soup = BeautifulSoup(html, "html.parser")
    nfo = soup.find("div", class_=re.compile(r"topic-nfo"))
    if nfo is None:
        return ""
    label = nfo.find("strong")
    if label and "簡介" in label.get_text():
        label.decompose()
    for img in nfo.find_all("img"):
        src = img.get("src", "")
        if src.startswith("//"):
            img["src"] = "https:" + src
        elif src.startswith("/"):
            img["src"] = BASE_URL + src
    for a in nfo.find_all("a"):
        href = a.get("href", "")
        if href.startswith("/"):
            a["href"] = BASE_URL + href
    return nfo.decode_contents().strip()


def fetch_pages(max_pages: int) -> list[dict]:
    """从第 1 页翻到最后一页，返回所有条目（按最新→最旧）。"""
    all_topics = []
    seen_ids = set()
    for page in range(1, max_pages + 1):
        url = page_url(page)
        print(f"    列表第 {page} 页: {url}")
        try:
            html = fetch_html(url)
        except Exception as e:
            print(f"      !! 抓取失败: {e}")
            break
        soup = BeautifulSoup(html, "html.parser")
        topics = parse_list_page(html)
        if not topics:
            break
        new_count = 0
        for t in topics:
            if t["id"] not in seen_ids:
                seen_ids.add(t["id"])
                all_topics.append(t)
                new_count += 1
        print(f"      解析 {len(topics)} 条，新增 {new_count} 条")
        if not has_next_page(soup):
            break
        if page < max_pages:
            time.sleep(PAGE_DELAY)
    return all_topics


def load_state(path: str) -> dict:
    """读取历史条目 {id: topic_dict}。兼容旧版 list 格式。"""
    if not os.path.exists(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            return {str(i): {"id": str(i)} for i in data}
        return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def save_state(path: str, state: dict) -> None:
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=2)


def to_rfc822(iso_str: str | None) -> str:
    if not iso_str:
        return datetime.now(TZ_CST).strftime("%a, %d %b %Y %H:%M:%S %z")
    try:
        dt = datetime.fromisoformat(iso_str)
    except ValueError:
        return datetime.now(TZ_CST).strftime("%a, %d %b %Y %H:%M:%S %z")
    return dt.strftime("%a, %d %b %Y %H:%M:%S %z")


def build_feed(topics: list[dict], out_path: str, feed_url: str = "") -> None:
    fg = FeedGenerator()
    fg.title(safe_xml(FEED_TITLE))
    fg.link(href=FEED_LINK, rel="alternate")
    self_url = feed_url or "file:///" + out_path.replace("\\", "/")
    fg.link(href=self_url, rel="self")
    fg.description(safe_xml(FEED_DESC))
    fg.language("zh-cn")

    def sort_key(t):
        if t.get("pub_date"):
            try:
                return datetime.fromisoformat(t["pub_date"])
            except ValueError:
                pass
        return datetime.min.replace(tzinfo=TZ_CST)

    for t in sorted(topics, key=sort_key):  # feedgen add_entry 头插，升序输入=最新在前
        fe = fg.add_entry()
        fe.title(safe_xml(t["title"]))
        fe.link(href=t["link"])
        fe.guid(t["link"], permalink=True)
        fe.pubDate(to_rfc822(t.get("pub_date")))
        fe.category({"term": safe_xml(t.get("category", ""))})
        parts = []
        if t.get("size"):
            parts.append(f"大小：{safe_xml(t['size'])}")
        if t.get("publisher"):
            parts.append(f"发布人：{safe_xml(t['publisher'])}")
        if t.get("magnet"):
            parts.append(f'<p>磁力链接：<a href="{t["magnet"]}">magnet</a></p>')
        meta = "<br/>".join(parts)
        content = safe_xml(t.get("content") or "")
        if content:
            fe.description(meta + "<hr/>" + content)
            fe.content(content, type="html")
        else:
            fe.description(meta)

    fg.rss_file(out_path, pretty=True)


def export_xlsx(topics: list[dict], xlsx_path: str) -> None:
    """把全部条目导出成带样式、带超链接的 Excel 清单。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "VCB-Studio 发布清单"

    headers = ["发布日期", "标题", "分类", "大小", "发布人"]
    ws.append(headers)
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="17365D")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    zebra = PatternFill("solid", fgColor="DCE6F1")

    def sort_key(t):
        if t.get("pub_date"):
            try:
                return datetime.fromisoformat(t["pub_date"])
            except ValueError:
                pass
        return datetime.min.replace(tzinfo=TZ_CST)

    for row_num, t in enumerate(sorted(topics, key=sort_key, reverse=True), start=2):
        pub = t.get("pub_date") or ""
        if pub:
            try:
                pub = datetime.fromisoformat(pub).strftime("%Y-%m-%d %H:%M")
            except ValueError:
                pass
        title = t.get("title", "")
        if title.startswith(("=", "+", "-", "@")):
            title = "'" + title
        pub = safe_xml(pub)
        title = safe_xml(title)
        ws.cell(row=row_num, column=1, value=pub)
        c_title = ws.cell(row=row_num, column=2, value=title)
        link = t.get("link", "")
        if link:
            c_title.hyperlink = link
            c_title.style = "Hyperlink"
        ws.cell(row=row_num, column=3, value=safe_xml(t.get("category", "")))
        ws.cell(row=row_num, column=4, value=safe_xml(t.get("size", "")))
        ws.cell(row=row_num, column=5, value=safe_xml(t.get("publisher", "")))
        if row_num % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_num, column=col).fill = zebra

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 12
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(1, len(topics) + 1)}"
    wb.save(xlsx_path)


def main() -> int:
    ap = argparse.ArgumentParser(description="DMHY RSS 生成器")
    ap.add_argument("--force", action="store_true", help="忽略去重状态，重新抓取并生成")
    ap.add_argument("--out", default=DEFAULT_OUT, help="输出 XML 文件路径")
    ap.add_argument("--feed-url", default="", help="feed 的公开订阅地址（self link）")
    ap.add_argument("--max-pages", type=int, default=MAX_PAGES, help=f"列表翻页安全上限（默认 {MAX_PAGES}）")
    ap.add_argument("--detail-batch", type=int, default=DETAIL_BATCH, help=f"每次补爬历史详情条数（默认 {DETAIL_BATCH}）")
    ap.add_argument("--no-detail", action="store_true", help="不抓取详情页正文")
    args = ap.parse_args()

    print(f"[1/5] 抓取列表全部页（上限 {args.max_pages}）...")
    fresh_topics = fetch_pages(args.max_pages)
    if not fresh_topics:
        print("!! 未解析到任何条目")
        return 1
    print(f"    本次列表共 {len(fresh_topics)} 条")

    state = load_state(STATE_FILE)
    if args.force:
        new_items = fresh_topics
    else:
        new_items = [t for t in fresh_topics if t["id"] not in state]
    print(f"[2/5] 新增条目 {len(new_items)} 条（历史已有 {len(state)} 条）")

    for t in new_items:
        state[t["id"]] = t

    if FETCH_DETAIL and not args.no_detail and new_items:
        print(f"[3/5] 抓取 {len(new_items)} 条新条目详情正文...")
        for i, t in enumerate(new_items, 1):
            t["content"] = fetch_detail_content(t["link"])
            if i % 20 == 0 or i == len(new_items):
                print(f"      {i}/{len(new_items)}")
            time.sleep(DETAIL_DELAY)
    else:
        print("[3/5] 跳过新条目详情")

    if FETCH_DETAIL and not args.no_detail:
        missing = [t for t in state.values() if not t.get("content")]
        missing.sort(key=lambda t: t.get("pub_date") or "", reverse=True)
        backfill = missing[: args.detail_batch]
        if backfill:
            print(f"[4/5] 历史回填：{len(missing)} 条无正文，本次补 {len(backfill)} 条...")
            for i, t in enumerate(backfill, 1):
                t["content"] = fetch_detail_content(t["link"])
                if i % 20 == 0 or i == len(backfill):
                    print(f"      {i}/{len(backfill)}")
                time.sleep(DETAIL_DELAY)
        else:
            print("[4/5] 所有条目均已有正文，无需回填")
    else:
        print("[4/5] 跳过历史回填")

    all_topics = list(state.values())
    build_feed(all_topics, args.out, args.feed_url)
    print(f"[5/5] 已生成 feed: {args.out}（共 {len(all_topics)} 条）")
    try:
        export_xlsx(all_topics, XLSX_OUT)
        print(f"      已生成清单: {XLSX_OUT}")
    except Exception as e:
        print(f"      !! xlsx 生成失败: {e}")

    save_state(STATE_FILE, state)
    return 0


if __name__ == "__main__":
    sys.exit(main())
